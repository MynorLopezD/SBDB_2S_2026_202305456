"""
generar_import_sql.py

Genera el script SQL de importacion (INSERT INTO ...) para el modelo de la
Practica 2 (Programa de EPS) a partir del archivo Excel del dataset.

Requisito: la libreria openpyxl.
    pip install openpyxl

Uso:
    python generar_import_sql.py
        Busca "Dataset_Practica2.xlsx" en la misma carpeta y genera
        "Import_Practica2.sql" en la misma carpeta.

    python generar_import_sql.py ruta/al/excel.xlsx ruta/al/salida.sql
        Usa rutas personalizadas de entrada y salida.

Como usarlo para modificar/agregar datos:
    1. Edita el archivo Excel (agrega filas, cambia valores, etc.)
       respetando las mismas hojas y columnas que ya tiene.
    2. Guarda el Excel.
    3. Vuelve a correr este script.
    4. Se genera un Import_Practica2.sql actualizado, listo para
       ejecutar con "Ejecutar Script" (F5) en Oracle SQL Developer.

Notas de las transformaciones aplicadas (mismas reglas usadas
originalmente en la Practica 2):
    - especialidad: catalogo propio, construido con los valores UNICOS
      de PLAZA.ESPECIALIDAD_TECNICA + CATEDRATICO.ESPECIALIDAD +
      ESTUDIANTE.CARRERA_TECNICA, en el orden en que aparecen.
    - estudiante.primera_practica = 1 - ES_REPITENCIA
    - colocacion.activo = '1' si el ID_ESTADO corresponde al estado
      llamado 'Activa' en la hoja ESTADO_COLOCACION, si no '0'.
"""

import sys
import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit(
        "Falta la libreria openpyxl. Instalala con:\n"
        "    pip install openpyxl"
    )


# ------------------------------------------------------------------
# Utilidades de lectura y formateo
# ------------------------------------------------------------------

def leer_filas(wb, hoja):
    """Devuelve las filas de datos (sin encabezado, sin filas vacias)."""
    ws = wb[hoja]
    todas = list(ws.iter_rows(values_only=True))
    return [r for r in todas[1:] if r[0] is not None]


def esc(valor):
    """Escapa comillas simples para SQL."""
    return str(valor).replace("'", "''")


def sqlstr(valor):
    if valor is None or valor == "":
        return "NULL"
    return f"'{esc(valor)}'"


def sqlnum(valor):
    if valor is None or valor == "":
        return "NULL"
    f = float(valor)
    return str(int(f)) if f == int(f) else str(f)


def sqldate(valor):
    if valor is None or valor == "":
        return "NULL"
    if isinstance(valor, datetime.datetime):
        return f"TO_DATE('{valor.strftime('%Y-%m-%d')}','YYYY-MM-DD')"
    return "NULL"


# ------------------------------------------------------------------
# Generacion del script
# ------------------------------------------------------------------

def generar_sql(ruta_excel: Path) -> str:
    wb = load_workbook(ruta_excel, read_only=True, data_only=True)
    out = []

    out.append("-- ============================================================")
    out.append("-- PRACTICA 2 - Script de importacion de datos")
    out.append(f"-- Generado automaticamente desde: {ruta_excel.name}")
    out.append(f"-- Fecha de generacion: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    out.append("-- Ejecutar como Script (F5) sobre la conexion de Oracle correspondiente.")
    out.append("-- Respeta el orden de dependencia de llaves foraneas.")
    out.append("-- ============================================================")
    out.append("")

    # 1. SECTOR_ECONOMICO
    out.append("-- 1. sector_economico")
    for r in leer_filas(wb, "SECTOR_ECONOMICO"):
        out.append(
            f"INSERT INTO sector_economico (id_sector, nombre) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])});"
        )
    out.append("")

    # 2. DEPARTAMENTO
    out.append("-- 2. departamento")
    for r in leer_filas(wb, "DEPARTAMENTO"):
        out.append(
            f"INSERT INTO departamento (id_departamento, nombre) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])});"
        )
    out.append("")

    # 3. ESTADO_COLOCACION
    estado_rows = leer_filas(wb, "ESTADO_COLOCACION")
    out.append("-- 3. estado_colocacion")
    for r in estado_rows:
        out.append(
            f"INSERT INTO estado_colocacion (id_estado, nombre) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])});"
        )
    out.append("")

    # 4. TIPO_EVALUACION
    out.append("-- 4. tipo_evaluacion")
    for r in leer_filas(wb, "TIPO_EVALUACION"):
        out.append(
            f"INSERT INTO tipo_evaluacion (id_tipo_evaluacion, nombre) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])});"
        )
    out.append("")

    # 5. CRITERIO
    out.append("-- 5. criterio")
    for r in leer_filas(wb, "CRITERIO"):
        out.append(
            f"INSERT INTO criterio (id_criterio, nombre) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])});"
        )
    out.append("")

    # 6. INSTITUTO
    out.append("-- 6. instituto")
    for r in leer_filas(wb, "INSTITUTO"):
        out.append(
            f"INSERT INTO instituto (id_instituto, nombre, direccion, codigo_autorizacion) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])}, {sqlstr(r[2])}, {sqlstr(r[3])});"
        )
    out.append("")

    # 7. ESPECIALIDAD (catalogo propio construido a partir de 3 hojas)
    plaza_rows = leer_filas(wb, "PLAZA")
    cated_rows = leer_filas(wb, "CATEDRATICO")
    est_rows = leer_filas(wb, "ESTUDIANTE")

    especialidades = []
    vistos = set()
    for r in plaza_rows:
        if r[1] not in vistos:
            vistos.add(r[1])
            especialidades.append(r[1])
    for r in cated_rows:
        if r[4] not in vistos:
            vistos.add(r[4])
            especialidades.append(r[4])
    for r in est_rows:
        if r[2] not in vistos:
            vistos.add(r[2])
            especialidades.append(r[2])

    esp_id = {nombre: i + 1 for i, nombre in enumerate(especialidades)}

    out.append("-- 7. especialidad (catalogo propio: valores unicos tomados de")
    out.append("--    PLAZA.ESPECIALIDAD_TECNICA + CATEDRATICO.ESPECIALIDAD + ESTUDIANTE.CARRERA_TECNICA)")
    for nombre, eid in esp_id.items():
        out.append(f"INSERT INTO especialidad (id_especialidad, nombre) VALUES ({eid}, {sqlstr(nombre)});")
    out.append("")

    # 8. MUNICIPIO
    out.append("-- 8. municipio")
    for r in leer_filas(wb, "MUNICIPIO"):
        out.append(
            f"INSERT INTO municipio (id_municipio, nombre, departamento_id) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])}, {sqlnum(r[2])});"
        )
    out.append("")

    # 9. EMPRESA
    out.append("-- 9. empresa")
    for r in leer_filas(wb, "EMPRESA"):
        out.append(
            f"INSERT INTO empresa (id_empresa, nombre, direccion, sector_id) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])}, {sqlstr(r[2])}, {sqlnum(r[3])});"
        )
    out.append("")

    # 10. CONTACTO
    out.append("-- 10. contacto")
    for r in leer_filas(wb, "CONTACTO_EMPRESARIAL"):
        out.append(
            f"INSERT INTO contacto (id_contacto, nombre, telefono, correo, empresa_id) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(r[1])}, {sqlstr(int(r[2]))}, {sqlstr(r[3])}, {sqlnum(r[4])});"
        )
    out.append("")

    # 11. CATEDRATICO
    out.append("-- 11. catedratico (especialidad_id resuelto via el catalogo generado arriba)")
    for r in cated_rows:
        eid = esp_id[r[4]]
        out.append(
            f"INSERT INTO catedratico (id_catedratico, identificacion, nombre, telefono, especialidad_id, instituto_id) "
            f"VALUES ({sqlnum(r[0])}, {sqlstr(int(r[1]))}, {sqlstr(r[2])}, {sqlstr(int(r[3]))}, {eid}, {sqlnum(r[5])});"
        )
    out.append("")

    # 12. ESTUDIANTE
    out.append("-- 12. estudiante (PK=carnet, especialidad_id resuelto, primera_practica = 1 - es_repitencia)")
    for r in est_rows:
        carnet, nombre_completo, carrera, direccion, telefono, fecha_nac, genero, es_repit, id_mun, id_inst = r
        eid = esp_id[carrera]
        primera = 0 if int(es_repit) == 1 else 1
        out.append(
            "INSERT INTO estudiante (carnet, nombre_completo, direccion, telefono, fecha_nacimiento, genero, "
            "primera_practica, especialidad_id, municipio_id, instituto_id) VALUES "
            f"({sqlnum(carnet)}, {sqlstr(nombre_completo)}, {sqlstr(direccion)}, {sqlstr(int(telefono))}, "
            f"{sqldate(fecha_nac)}, {sqlstr(genero)}, {primera}, {eid}, {sqlnum(id_mun)}, {sqlnum(id_inst)});"
        )
    out.append("")

    # 13. PLAZA
    out.append("-- 13. plaza (especialidad_id resuelto)")
    for r in plaza_rows:
        id_plaza, esp_txt, id_emp, id_con = r
        eid = esp_id[esp_txt]
        out.append(
            f"INSERT INTO plaza (id_plaza, especialidad_id, empresa_id, contacto_id) "
            f"VALUES ({sqlnum(id_plaza)}, {eid}, {sqlnum(id_emp)}, {sqlnum(id_con)});"
        )
    out.append("")

    # 14. COLOCACION
    activa_id = next(r[0] for r in estado_rows if r[1] == "Activa")
    out.append("-- 14. colocacion (activo derivado: '1' si estado = Activa, '0' en otro caso)")
    for r in leer_filas(wb, "COLOCACION"):
        id_col, f_ini, f_fin, id_est, id_pla, id_cat, id_estado = r
        activo = "'1'" if int(id_estado) == int(activa_id) else "'0'"
        out.append(
            "INSERT INTO colocacion (id_colocacion, fecha_inicio, fecha_fin, estudiante_id, plaza_id, "
            "catedratico_id, estado_id, activo) VALUES "
            f"({sqlnum(id_col)}, {sqldate(f_ini)}, {sqldate(f_fin)}, {sqlnum(id_est)}, {sqlnum(id_pla)}, "
            f"{sqlnum(id_cat)}, {sqlnum(id_estado)}, {activo});"
        )
    out.append("")

    # 15. BITACORA
    out.append("-- 15. bitacora")
    for r in leer_filas(wb, "BITACORA"):
        id_bit, correl, fecha, horas, act, obs, id_col, id_contval = r
        out.append(
            "INSERT INTO bitacora (id_bitacora, correlativo, fecha, horas_trabajadas, actividades_realizadas, "
            "observaciones, colocacion_id, validado_por) VALUES "
            f"({sqlnum(id_bit)}, {sqlnum(correl)}, {sqldate(fecha)}, {sqlnum(horas)}, {sqlstr(act)}, "
            f"{sqlstr(obs)}, {sqlnum(id_col)}, {sqlnum(id_contval)});"
        )
    out.append("")

    # 16. EVALUACION
    out.append("-- 16. evaluacion")
    for r in leer_filas(wb, "EVALUACION"):
        id_ev, fecha, id_col, id_cat, id_tipo = r
        out.append(
            "INSERT INTO evaluacion (id_evaluacion, fecha, colocacion_id, catedratico_id, tipo_evaluacion_id) "
            f"VALUES ({sqlnum(id_ev)}, {sqldate(fecha)}, {sqlnum(id_col)}, {sqlnum(id_cat)}, {sqlnum(id_tipo)});"
        )
    out.append("")

    # 17. EVALUACION_CRITERIO
    out.append("-- 17. evaluacion_criterio")
    for r in leer_filas(wb, "DETALLE_EVALUACION"):
        id_ev, id_crit, punt = r
        out.append(
            f"INSERT INTO evaluacion_criterio (evaluacion_id, criterio_id, puntaje) "
            f"VALUES ({sqlnum(id_ev)}, {sqlnum(id_crit)}, {sqlnum(punt)});"
        )
    out.append("")

    out.append("COMMIT;")

    return "\n".join(out)


# ------------------------------------------------------------------
# Punto de entrada
# ------------------------------------------------------------------

def main():
    if len(sys.argv) >= 3:
        ruta_excel = Path(sys.argv[1])
        ruta_salida = Path(sys.argv[2])
    elif len(sys.argv) == 2:
        ruta_excel = Path(sys.argv[1])
        ruta_salida = Path("Import_Practica2.sql")
    else:
        ruta_excel = Path("Dataset_Practica2.xlsx")
        ruta_salida = Path("Import_Practica2.sql")

    if not ruta_excel.exists():
        sys.exit(f"No se encontro el archivo Excel: {ruta_excel}")

    print(f"Leyendo: {ruta_excel}")
    sql = generar_sql(ruta_excel)

    ruta_salida.write_text(sql, encoding="utf-8")
    print(f"Listo. Script generado en: {ruta_salida}")
    print("Ejecutalo en Oracle SQL Developer con 'Ejecutar Script' (F5).")


if __name__ == "__main__":
    main()